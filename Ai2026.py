import io
import json
import os
import streamlit as st
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt

# Google GenAI SDK
try:
    from google import genai
    from google.genai import types
    HAS_GENAI = True
except ImportError:
    HAS_GENAI = False

# ---------------------------------------------------------
# 1. CẤU HÌNH TRANG STREAMLIT
# ---------------------------------------------------------
st.set_page_config(
    page_title="Nova AI - Trợ lý & Tạo Slide Thông Minh",
    page_icon="✨",
    layout="centered",
    initial_sidebar_state="expanded"
)

# ---------------------------------------------------------
# 2. CSS TÙY CHỈNH (Giao diện hiện đại & Chuẩn Di Động)
# ---------------------------------------------------------
st.markdown("""
<style>
    /* Ẩn bớt thành phần thừa của Streamlit */
    header { visibility: hidden; }
    footer { visibility: hidden; }
    #MainMenu { visibility: hidden; }
    
    .block-container {
        padding-top: 1.5rem !important;
        padding-bottom: 4rem !important;
        max-width: 720px;
    }

    .hero-title {
        text-align: center;
        font-size: 2.2rem;
        font-weight: 800;
        color: #0f172a;
        margin-top: 10px;
        margin-bottom: 8px;
        line-height: 1.2;
    }
    
    .hero-sub {
        text-align: center;
        color: #64748b;
        font-size: 0.95rem;
        margin-bottom: 1.8rem;
        line-height: 1.5;
    }

    div.stButton > button {
        width: 100%;
        border-radius: 12px;
        border: 1px solid #e2e8f0;
        padding: 12px 16px;
        background-color: #ffffff;
        box-shadow: 0 2px 4px rgba(0,0,0,0.02);
        transition: all 0.2s ease;
    }
    
    div.stButton > button:hover {
        border-color: #6366f1;
        background-color: #f8fafc;
    }
</style>
""", unsafe_allow_html=True)

# ---------------------------------------------------------
# 3. PHẦN CÀI ĐẶT API KEY (SIDEBAR)
# ---------------------------------------------------------
with st.sidebar:
    st.title("⚙️ Cấu hình AI")
    
    # Lấy API Key từ Streamlit Secrets hoặc môi trường nếu có
    env_api_key = os.environ.get("GEMINI_API_KEY", "")
    if "GEMINI_API_KEY" in st.secrets:
        env_api_key = st.secrets["GEMINI_API_KEY"]

    api_key_input = st.text_input(
        "Nhập Google Gemini API Key:",
        value=env_api_key,
        type="password",
        help="Lấy API Key miễn phí tại https://aistudio.google.com/"
    )
    
    st.markdown("""
    ---
    💡 **Hướng dẫn lấy API Key miễn phí:**
    1. Truy cập [Google AI Studio](https://aistudio.google.com/)
    2. Đăng nhập Google & bấm **Get API key**
    3. Dán mã Key vào ô trên để kích hoạt AI thật!
    """)

# Khởi tạo Gemini Client nếu có API Key
client = None
if api_key_input.strip() and HAS_GENAI:
    try:
        client = genai.Client(api_key=api_key_input.strip())
    except Exception as e:
        st.sidebar.error(f"Lỗi khởi tạo API: {e}")

# ---------------------------------------------------------
# 4. HÀM TẠO FILE XUẤT (PPTX, DOCX, PYTHON)
# ---------------------------------------------------------
def generate_pptx(topic, slides_data):
    """Tạo tệp PowerPoint (.pptx) trong bộ nhớ RAM"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide Tiêu đề
    blank_layout = prs.slide_layouts[6]
    title_slide = prs.slides.add_slide(blank_layout)
    tx_box = title_slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(2))
    tf = tx_box.text_frame
    p = tf.paragraphs[0]
    p.text = topic
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.name = "Arial"

    # Slide Nội dung
    for slide_item in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title_shape = slide.shapes.title
        title_shape.text = slide_item.get("title", "Slide")
        
        body_shape = slide.placeholders[1]
        tf_body = body_shape.text_frame
        tf_body.word_wrap = True
        
        bullets = slide_item.get("bullets", [])
        for idx, bullet in enumerate(bullets):
            if idx == 0:
                p_item = tf_body.paragraphs[0]
                p_item.text = bullet
            else:
                p_item = tf_body.add_paragraph()
                p_item.text = bullet

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer


def generate_docx(topic, slides_data):
    """Tạo tệp Word (.docx) trong bộ nhớ RAM"""
    doc = Document()
    doc.add_heading(topic, 0)

    for slide in slides_data:
        doc.add_heading(slide.get("title", "Slide"), level=1)
        for bullet in slide.get("bullets", []):
            doc.add_paragraph(bullet, style='List Bullet')

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def generate_python_script(topic, slides_data):
    """Tạo mã nguồn Python (.py) có thể chạy độc lập"""
    code = f'# -*- coding: utf-8 -*-\n'
    code += f'"""\nNova AI Generated Presentation Script\nTopic: {topic}\n"""\n\n'
    code += f'topic = "{topic}"\n\n'
    code += 'slides = [
'
    
    for slide in slides_data:
        code += '    {
'
        code += f'        "title": "{slide.get("title", "")}",
'
        code += '        "bullets": [
'
        for bullet in slide.get("bullets", []):
            code += f'            "{bullet}",
'
        code += '        ]
'
        code += '    },
'
        
    code += ']

'
    code += 'def main():
'
    code += '    print(f"=== BÀI THUYẾT TRÌNH: {topic} ===")
'
    code += '    for idx, slide in enumerate(slides, 1):
'
    code += '        print(f"\n[Slide {idx}] {slide['title']}")
'
    code += '        for item in slide["bullets"]:
'
    code += '            print(f"  - {item}")

'
    code += 'if __name__ == "__main__":
'
    code += '    main()
'

    return code.encode("utf-8")

# ---------------------------------------------------------
# 5. INITIALIZE SESSION STATE
# ---------------------------------------------------------
if "mode" not in st.session_state:
    st.session_state.mode = "presentation"

if "messages" not in st.session_state:
    st.session_state.messages = []

if "generated_slides" not in st.session_state:
    st.session_state.generated_slides = None

if "current_topic" not in st.session_state:
    st.session_state.current_topic = ""

# ---------------------------------------------------------
# 6. HEADER VÀ NÚT CHUYỂN CHẾ ĐỘ
# ---------------------------------------------------------
st.markdown('<div class="hero-title">What can I help you create?</div>', unsafe_allow_html=True)
st.markdown('<div class="hero-sub">Hỏi AI bất cứ điều gì — hoặc yêu cầu tạo bài trình bày và xuất file PowerPoint, Word, Python.</div>', unsafe_allow_html=True)

col1, col2 = st.columns(2)

with col1:
    if st.button("💬  **Hỏi đáp AI**

Hỏi đáp trực tiếp cùng mô hình Gemini AI."):
        st.session_state.mode = "chat"

with col2:
    if st.button("📊  **Tạo trình bày**

Tạo slide thuyết trình AI & Xuất tệp PPTX/Word/Python."):
        st.session_state.mode = "presentation"

st.divider()

# ---------------------------------------------------------
# 7. CHẾ ĐỘ 1: TẠO BÀI TRÌNH BÀY VỚI GEMINI AI
# ---------------------------------------------------------
if st.session_state.mode == "presentation":
    st.subheader("📊 Tạo bài thuyết trình thông minh bằng AI")
    
    topic_input = st.text_input(
        "Chủ đề bài thuyết trình:", 
        placeholder="Ví dụ: Ứng dụng AI trong giáo dục năm 2026..."
    )
    
    num_slides = st.slider("Số lượng Slide:", min_value=3, max_value=10, value=4)

    if st.button("🚀 Bắt đầu tạo Slide bằng AI", type="primary"):
        if not topic_input.strip():
            st.warning("Vui lòng nhập chủ đề bài thuyết trình.")
        else:
            st.session_state.current_topic = topic_input.strip()
            
            # Nếu có API Key -> Gọi Gemini AI thật
            if client:
                with st.spinner("AI đang tư duy và lập dàn ý slide..."):
                    try:
                        prompt = f"""Bạn là một chuyên gia thuyết trình. Hãy tạo bài thuyết trình gồm {num_slides} slide cho chủ đề: "{topic_input}".
Yêu cầu trả về đúng định dạng JSON dạng danh sách (list) các object. Mỗi object gồm:
- "title": Tiêu đề của slide (ngắn gọn, hấp dẫn)
- "bullets": Danh sách 3-4 ý chính dạng chuỗi văn bản.

Ví dụ định dạng JSON yêu cầu:
[
  {{"title": "1. Giới thiệu", "bullets": ["Ý 1", "Ý 2", "Ý 3"]}},
  {{"title": "2. Thách thức", "bullets": ["Ý 1", "Ý 2", "Ý 3"]}}
]
Chỉ trả về chuỗi JSON thuần túy, không kèm mã markdown codeblock hay văn bản khác.
"""
                        response = client.models.generate_content(
                            model="gemini-2.5-flash",
                            contents=prompt
                        )
                        
                        raw_text = response.text.strip()
                        # Làm sạch nếu AI trả về markdown code block ```json ... ```
                        if raw_text.startswith("```"):
                            raw_text = raw_text.split("```")[1]
                            if raw_text.startswith("json"):
                                raw_text = raw_text[4:]
                        raw_text = raw_text.strip()
                        
                        slides_json = json.loads(raw_text)
                        st.session_state.generated_slides = slides_json
                        st.success("🎉 Gemini AI đã khởi tạo thành công bài thuyết trình!")
                    except Exception as e:
                        st.error(f"Lỗi khi gọi Gemini AI: {e}. Đang dùng chế độ tạo dàn ý dự phòng.")
                        # Dàn ý dự phòng nếu lỗi
                        st.session_state.generated_slides = [
                            {"title": f"1. Tổng quan về {topic_input}", "bullets": ["Khái niệm cơ bản.", "Lý do chủ đề quan trọng hiện nay.", "Mục tiêu bài thuyết trình."]},
                            {"title": "2. Phân tích chi tiết", "bullets": ["Các khía cạnh cốt lõi.", "Ưu điểm và nhược điểm.", "Ví dụ thực tế."]},
                            {"title": "3. Định hướng & Kết luận", "bullets": ["Bài học rút ra.", "Khuyên dùng & Giải pháp.", "Q&A giải đáp."]}
                        ]
            else:
                # Nếu chưa nhập API Key -> Thông báo & Tạo slide mẫu
                st.info("💡 Bạn chưa nhập Gemini API Key ở thanh bên trái (Sidebar). Hệ thống đang tạo bản mẫu thử nghiệm.")
                st.session_state.generated_slides = [
                    {"title": f"1. Giới thiệu về {topic_input}", "bullets": ["Khái niệm tổng quan.", "Xu hướng phát triển hiện tại.", "Tầm quan trọng của chủ đề."]},
                    {"title": "2. Các giá trị cốt lõi", "bullets": ["Tăng hiệu suất công việc.", "Tối ưu hóa quy trình.", "Đổi mới sáng tạo."]},
                    {"title": "3. Kết luận & Q&A", "bullets": ["Tóm tắt nội dung chính.", "Đề xuất hành động tiếp theo.", "Giải đáp thắc mắc."]}
                ][:num_slides]
                st.success("Đã tạo xong khung slide thử nghiệm!")

    # Hiển thị kết quả & Nút xuất file
    if st.session_state.generated_slides:
        st.write("---")
        st.markdown(f"### 📋 Xem trước nội dung: **{st.session_state.current_topic}**")

        for idx, slide in enumerate(st.session_state.generated_slides, 1):
            with st.expander(f"Slide {idx}: {slide.get('title', '')}", expanded=True):
                for bullet in slide.get("bullets", []):
                    st.write(f"• {bullet}")

        st.markdown("### 📥 Xuất Tệp Trình Bày")
        st.caption("Tải về tệp hoàn chỉnh để sử dụng ngay:")

        exp_col1, exp_col2, exp_col3 = st.columns(3)

        with exp_col1:
            pptx_file = generate_pptx(st.session_state.current_topic, st.session_state.generated_slides)
            st.download_button(
                label="📄 PowerPoint (.pptx)",
                data=pptx_file,
                file_name=f"{st.session_state.current_topic}.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                use_container_width=True
            )

        with exp_col2:
            docx_file = generate_docx(st.session_state.current_topic, st.session_state.generated_slides)
            st.download_button(
                label="📝 Word (.docx)",
                data=docx_file,
                file_name=f"{st.session_state.current_topic}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )

        with exp_col3:
            py_file = generate_python_script(st.session_state.current_topic, st.session_state.generated_slides)
            st.download_button(
                label="🐍 Python Script (.py)",
                data=py_file,
                file_name=f"{st.session_state.current_topic}.py",
                mime="text/x-python",
                use_container_width=True
            )

# ---------------------------------------------------------
# 8. CHẾ ĐỘ 2: HỎI ĐÁP AI (CHATBOT)
# ---------------------------------------------------------
elif st.session_state.mode == "chat":
    st.subheader("💬 Trò chuyện trực tiếp cùng Gemini AI")

    # Hiển thị lịch sử chat
    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    if prompt := st.chat_input("Hỏi AI bất cứ điều gì..."):
        st.session_state.messages.append({"role": "user", "content": prompt})
        with st.chat_message("user"):
            st.markdown(prompt)

        with st.chat_message("assistant"):
            if client:
                with st.spinner("Gemini AI đang trả lời..."):
                    try:
                        response = client.models.generate_content(
                            model="gemini-2.5-flash",
                            contents=prompt
                        )
                        reply_text = response.text
                    except Exception as e:
                        reply_text = f"Lỗi gọi AI: {e}"
            else:
                reply_text = f"🤖 [Chế độ thử nghiệm] Tôi đã nhận được câu hỏi của bạn: **'{prompt}'**.

*Mẹo: Hãy nhập Gemini API Key ở menu bên trái (Sidebar) để kích hoạt câu trả lời thông minh thật từ Gemini AI nhé!*"

            st.markdown(reply_text)
            st.session_state.messages.append({"role": "assistant", "content": reply_text})
